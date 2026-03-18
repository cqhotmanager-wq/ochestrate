"""文件工具：支持多格式文件读写与目录枚举。"""

from __future__ import annotations

import json
import os
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from pypdf import PdfReader

from app.tools.security import ToolSecurityConfig


class FileTool:
    name = "file_tool"
    description = "Read, write, and list files including json/doc/docx/xlsx/pdf/txt/csv."
    required_roles = ["employee", "manager", "admin"]
    idempotent = False

    def __init__(self, security: ToolSecurityConfig, workspace: Path) -> None:
        self._security = security
        self._workspace = workspace.resolve()

    def run(self, params: dict[str, Any]) -> dict[str, Any]:
        operation = str(params.get("operation", "read")).lower()
        path_value = str(params.get("resource_path") or params.get("path") or "").strip()
        if not path_value:
            raise ValueError("resource_path is required for file tool")
        target = (self._workspace / Path(path_value)).resolve()
        if not self._security.is_path_allowed(target):
            raise PermissionError(f"path '{target}' is outside allowed directories")

        if operation == "list":
            if not target.exists() or not target.is_dir():
                raise ValueError("list operation requires an existing directory path")
            entries = [p.name for p in target.iterdir()]
            return {"operation": "list", "path": str(target), "entries": entries}

        if operation == "read":
            return self._read_file(target)
        if operation == "write":
            return self._write_file(target, params.get("content"), params.get("options", {}))
        raise ValueError(f"unsupported operation '{operation}'")

    def _read_file(self, target: Path) -> dict[str, Any]:
        self._guard_file_size(target)
        suffix = target.suffix.lower()
        if suffix == ".json":
            data = json.loads(target.read_text(encoding="utf-8"))
            return {"operation": "read", "resource_type": "json", "path": str(target), "data": data}
        if suffix in {".txt", ".csv"}:
            text = target.read_text(encoding="utf-8", errors="ignore")
            return {"operation": "read", "resource_type": suffix.lstrip("."), "path": str(target), "text": text}
        if suffix == ".docx":
            doc = Document(str(target))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return {"operation": "read", "resource_type": "docx", "path": str(target), "text": text}
        if suffix == ".doc":
            text = self._read_doc_via_conversion(target)
            return {"operation": "read", "resource_type": "doc", "path": str(target), "text": text}
        if suffix in {".xlsx", ".xls"}:
            frame = pd.read_excel(target)
            return {
                "operation": "read",
                "resource_type": "excel",
                "path": str(target),
                "rows": frame.to_dict(orient="records"),
            }
        if suffix == ".pdf":
            reader = PdfReader(str(target))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return {"operation": "read", "resource_type": "pdf", "path": str(target), "text": text}
        raise ValueError(f"unsupported file extension '{suffix}'")

    def _write_file(self, target: Path, content: Any, options: dict[str, Any]) -> dict[str, Any]:
        target.parent.mkdir(parents=True, exist_ok=True)
        suffix = target.suffix.lower()
        if suffix == ".json":
            target.write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
            return {"operation": "write", "resource_type": "json", "path": str(target), "status": "ok"}
        if suffix in {".txt", ".csv"}:
            text = content if isinstance(content, str) else json.dumps(content, ensure_ascii=False)
            target.write_text(text, encoding="utf-8")
            return {"operation": "write", "resource_type": suffix.lstrip("."), "path": str(target), "status": "ok"}
        if suffix == ".docx":
            doc = Document()
            if isinstance(content, list):
                for line in content:
                    doc.add_paragraph(str(line))
            else:
                doc.add_paragraph("" if content is None else str(content))
            doc.save(str(target))
            return {"operation": "write", "resource_type": "docx", "path": str(target), "status": "ok"}
        if suffix == ".doc":
            raise ValueError("writing '.doc' is not supported; write .docx instead")
        if suffix in {".xlsx", ".xls"}:
            if isinstance(content, list):
                frame = pd.DataFrame(content)
            elif isinstance(content, dict):
                frame = pd.DataFrame([content])
            else:
                frame = pd.DataFrame([{"value": content}])
            frame.to_excel(target, index=False)
            return {"operation": "write", "resource_type": "excel", "path": str(target), "status": "ok"}
        if suffix == ".pdf":
            raise ValueError("pdf write is not supported; export to txt or json")
        raise ValueError(f"unsupported file extension '{suffix}'")

    def _guard_file_size(self, target: Path) -> None:
        if not target.exists():
            raise FileNotFoundError(str(target))
        limit = self._security.max_file_size_mb * 1024 * 1024
        if os.path.getsize(target) > limit:
            raise ValueError(f"file exceeds max size limit ({self._security.max_file_size_mb}MB)")

    def _read_doc_via_conversion(self, source: Path) -> str:
        with tempfile.TemporaryDirectory() as tmp:
            cmd = [
                "soffice",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                tmp,
                str(source),
            ]
            try:
                subprocess.run(cmd, check=True, capture_output=True, text=True)
            except FileNotFoundError as exc:
                raise RuntimeError(
                    "'.doc' requires conversion tool (LibreOffice soffice). Install it and retry."
                ) from exc
            except subprocess.CalledProcessError as exc:
                raise RuntimeError(f"failed to convert .doc file: {exc.stderr}") from exc
            converted = Path(tmp) / f"{source.stem}.docx"
            if not converted.exists():
                raise RuntimeError("doc conversion finished but output file was not produced")
            doc = Document(str(converted))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())



